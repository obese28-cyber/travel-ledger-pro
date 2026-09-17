"""Read-only Excel snapshot of the customer account statement."""

from datetime import date, datetime, timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def build_customer_statement(statement):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Customer Account"
    customer, summary = statement["customer"], statement["summary"]

    def append(values):
        sheet.append(values)
        # Names, references and notes are text, never executable Excel formulas.
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.value = ILLEGAL_CHARACTERS_RE.sub("", cell.value)
                cell.data_type = "s"

    append(["CUSTOMER ACCOUNT STATEMENT"])
    append(["Customer", customer["name"]])
    append(["Customer ID", customer["id"]])
    append(["Email", customer.get("email"), "Phone", customer.get("phone")])
    append(["Exported (UTC)", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")])
    append(["Currency", "USD"])
    for label, key in [
        ("Total invoiced", "total_invoiced"),
        ("Applied payments", "applied_payments"),
        ("Outstanding balance", "net_outstanding"),
        ("Unused credit", "open_credit"),
        ("Total received (account summary)", "total_received"),
    ]:
        append([label, summary.get(key, 0)])
        sheet.cell(sheet.max_row, 2).number_format = '#,##0.00;[Red](#,##0.00)'
    append(["Snapshot of the current account. Credit applications are internal transfers, not new cash."])
    append(["Applied amounts and running balances match the account screen; unused credit is separate."])
    entries = statement["entries"]
    if entries and abs(entries[-1].get("running_balance", 0) - summary.get("net_outstanding", 0)) > 0.005:
        sheet["A13"] = "Review required: the account's final running balance differs from its outstanding summary."
        sheet["A13"].font = Font(bold=True, color="B45309")
    append([])
    append(["Date", "Type", "Reference", "Booking reference", "Description",
            "Charge (USD)", "Payment / credit amount (USD)", "Applied amount (USD)",
            "Unused credit (USD)", "Running balance (USD)", "DR / CR", "Status"])
    header_row = sheet.max_row
    for entry in entries:
        try:
            event_date = date.fromisoformat(entry.get("date") or "")
        except ValueError:
            event_date = entry.get("date") or ""
        append([
            event_date, entry["entry_type"].replace("_", " ").title(),
            entry.get("reference"), entry.get("booking_ref"), entry.get("description"),
            entry.get("debit", 0), entry.get("credit", 0), entry.get("amount_applied", 0),
            entry.get("unapplied_amount", 0), entry.get("running_balance", 0),
            entry.get("balance_label"),
            (entry.get("payment_status") or entry.get("status") or "").replace("_", " ").title(),
        ])
        sheet.cell(sheet.max_row, 1).number_format = "dd mmm yyyy"
        for column in range(6, 11):
            sheet.cell(sheet.max_row, column).number_format = '#,##0.00;[Red](#,##0.00)'
    sheet.freeze_panes = f"F{header_row + 1}"
    sheet.auto_filter.ref = f"A{header_row}:L{sheet.max_row}"
    for row in sheet:
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in sheet[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="243B53")
    sheet.row_dimensions[header_row].height = 44
    sheet["A1"].font = Font(size=18, bold=True, color="243B53")
    sheet.merge_cells("A1:L1")
    sheet.merge_cells("A12:L12")
    sheet.merge_cells("A13:L13")
    for column, width in enumerate([30, 28, 24, 26, 55, 18, 22, 20, 20, 22, 12, 24], 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = f"{header_row}:{header_row}"
    sheet.print_options.horizontalCentered = True
    sheet.print_area = sheet.dimensions
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
