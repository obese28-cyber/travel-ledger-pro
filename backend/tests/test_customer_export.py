"""Run with: python -m unittest discover -s tests"""
from io import BytesIO
import unittest

from flask import Flask
from flask_jwt_extended import create_access_token
from openpyxl import load_workbook

from app.extensions import db, jwt
from app.models import user, customer, booking, invoice, payment, vendor, vendor_bill, expense, accounting, audit, airline, trial_balance
from app.routes.customers import customers_bp


class CustomerExportTest(unittest.TestCase):
    def setUp(self):
        self.app = Flask(__name__)
        self.app.config.update(TESTING=True, SQLALCHEMY_DATABASE_URI="sqlite:///:memory:",
                               JWT_SECRET_KEY="test-secret-key-for-export-tests-only")
        db.init_app(self.app)
        jwt.init_app(self.app)
        self.app.register_blueprint(customers_bp, url_prefix="/api/customers")
        self.context = self.app.app_context()
        self.context.push()
        db.create_all()
        db.session.add(customer.Customer(id=1, name='=HYPERLINK("https://example.com")', phone="+1234"))
        db.session.commit()
        self.client = self.app.test_client()
        self.headers = {"Authorization": f"Bearer {create_access_token(identity='1')}"}

    def tearDown(self):
        db.session.remove()
        db.drop_all()
        self.context.pop()

    def export(self):
        response = self.client.get('/api/customers/1/statement?format=xlsx', headers=self.headers)
        self.assertEqual(response.status_code, 200)
        self.assertIn('attachment', response.headers['Content-Disposition'])
        return load_workbook(BytesIO(response.data)).active

    def test_empty_account_and_literal_text(self):
        sheet = self.export()
        self.assertEqual(sheet['B2'].data_type, 's')
        self.assertTrue(sheet['B2'].value.startswith('=HYPERLINK'))
        self.assertEqual(sheet['D4'].value, '+1234')
        self.assertEqual(sheet['B9'].value, 0)
        self.assertEqual(sheet['A15'].value, 'Date')
        self.assertEqual(sheet.freeze_panes, 'F16')

    def test_authentication_and_missing_customer(self):
        self.assertEqual(self.client.get('/api/customers/1/statement?format=xlsx').status_code, 401)
        self.assertEqual(self.client.get('/api/customers/999/statement?format=xlsx', headers=self.headers).status_code, 404)

    def test_partial_payment_and_applied_advance_match_statement(self):
        db.session.add(invoice.Invoice(id=1, invoice_number='INV-1', booking_id=1,
            customer_id=1, total_amount=100, amount_paid=60, status='partially_paid'))
        db.session.add_all([
            payment.Payment(payment_reference='PAY-1', customer_id=1, invoice_id=1,
                transaction_type='invoice_payment', amount=20, amount_applied=20,
                unapplied_amount=0, payment_method='cash'),
            payment.Payment(payment_reference='PAY-2', customer_id=1,
                transaction_type='advance_deposit', amount=70, amount_applied=40,
                unapplied_amount=30, payment_method='cash'),
            payment.Payment(payment_reference='PAY-3', customer_id=1, invoice_id=1,
                transaction_type='credit_application', amount=40, amount_applied=40,
                unapplied_amount=0, payment_method='credit_balance'),
        ])
        db.session.commit()
        data = self.client.get('/api/customers/1/statement', headers=self.headers).json['data']
        sheet = self.export()
        self.assertEqual(sheet['B9'].value, 40)
        self.assertEqual(sheet['B10'].value, 30)
        self.assertEqual(sheet['B11'].value, 90)  # internal transfer is not new cash
        for row, entry in enumerate(data['entries'], 16):
            self.assertEqual(sheet.cell(row, 10).value, entry['running_balance'])
            self.assertEqual(sheet.cell(row, 8).value, entry.get('amount_applied', 0))
        # Existing account discrepancy must be visible, never silently recalculated.
        self.assertIn('Review required', sheet['A13'].value)
        self.assertEqual(payment.Payment.query.count(), 3)


if __name__ == '__main__':
    unittest.main()
